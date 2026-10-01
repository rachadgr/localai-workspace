"""Tool-level tests: real artifact generation and real code execution."""

from __future__ import annotations

import os
import zipfile

from tools.registry import registry

registry.load_builtin_tools()


def test_registry_exposes_expected_tools():
    names = set(registry.names())
    expected = {"chat", "search", "research", "docs", "sheets", "slides", "developer", "website", "image", "files"}
    assert expected.issubset(names)


def test_docs_generates_real_docx_and_pdf(ctx):
    r = registry.execute("docs", ctx, {"title": "T Docx", "prompt": "A short brief about testing.", "format": "docx"})
    assert r.status == "SUCCESS", r.summary
    docx = [a for a in r.artifacts if a["name"].endswith(".docx")]
    assert docx and os.path.getsize(docx[0]["storage_path"]) > 0
    assert zipfile.is_zipfile(docx[0]["storage_path"])

    r2 = registry.execute("docs", ctx, {"title": "T Pdf", "prompt": "Another brief.", "format": "pdf"})
    pdf = [a for a in r2.artifacts if a["name"].endswith(".pdf")]
    assert pdf
    with open(pdf[0]["storage_path"], "rb") as fh:
        assert fh.read(5).startswith(b"%PDF")


def test_sheets_produces_valid_xlsx_with_statistics(ctx):
    rows = [{"item": f"I{i}", "qty": i, "price": i * 1.5} for i in range(1, 11)]
    r = registry.execute("sheets", ctx, {"title": "T Sheet", "rows": rows, "aggregate": [{"column": "qty", "op": "sum"}]})
    assert r.status == "SUCCESS", r.summary
    assert r.data["statistics"]["qty"]["sum"] == 55
    xlsx = [a for a in r.artifacts if a["name"].endswith(".xlsx")][0]
    assert zipfile.is_zipfile(xlsx["storage_path"])
    from openpyxl import load_workbook

    assert load_workbook(xlsx["storage_path"]).sheetnames


def test_slides_produces_pptx_with_8_to_12_slides(ctx):
    r = registry.execute("slides", ctx, {"title": "T Deck", "prompt": "A brief deck outline about testing.", "target_slides": 9})
    assert r.status == "SUCCESS", r.summary
    assert 8 <= r.data["slide_count"] <= 12
    pptx = [a for a in r.artifacts if a["name"].endswith(".pptx")][0]
    assert zipfile.is_zipfile(pptx["storage_path"])
    from pptx import Presentation

    assert len(Presentation(pptx["storage_path"]).slides) == r.data["slide_count"]


def test_website_builds_accessible_static_site(ctx):
    r = registry.execute("website", ctx, {"prompt": "Landing page for a coffee shop in Paris.", "site_name": "cafe-test"})
    assert r.status == "SUCCESS", r.summary
    validation = r.data["validation"]
    assert validation["valid"], validation
    html = r.data["files"]["index.html"]
    assert "<html lang=" in html and "viewport" in html and "<h1>" in html
    assert any(a["name"].endswith(".zip") for a in r.artifacts)


def test_developer_runs_real_tests_and_reports_truthfully(ctx):
    src = "def mul(a, b):\n    return a * b\n"
    test = "from main import mul\n\ndef test_mul():\n    assert mul(3, 4) == 12\n"
    r = registry.execute("developer", ctx, {"task": "verify mul", "language": "python", "files": {"main.py": src}, "tests": {"test_main.py": test}, "generate_tests": False})
    assert r.status == "SUCCESS"
    assert r.data["tests_executed"] is True
    assert r.data["tests_passed"] is True
    assert r.data["exit_code"] == 0


def test_developer_detects_failing_tests(ctx):
    src = "def mul(a, b):\n    return a + b\n"  # deliberately wrong
    test = "from main import mul\n\ndef test_mul():\n    assert mul(3, 4) == 12\n"
    r = registry.execute("developer", ctx, {"task": "verify mul", "language": "python", "files": {"main.py": src}, "tests": {"test_main.py": test}, "generate_tests": False, "max_fix_attempts": 0})
    assert r.data["tests_executed"] is True
    assert r.data["tests_passed"] is False


def test_image_generation_is_unavailable_without_provider(ctx):
    r = registry.execute("image", ctx, {"description": "a logo", "action": "generate"})
    assert r.status == "UNAVAILABLE"


def test_files_tool_blocks_traversal(ctx):
    r = registry.execute("files", ctx, {"action": "read", "path": "../../etc/passwd"})
    assert r.status == "FAILED"
