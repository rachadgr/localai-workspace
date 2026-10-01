"""AI Slides tool: generate real .pptx decks (8–12 slides default) with notes."""

from __future__ import annotations

import datetime as dt
import io
import json
import re
from typing import Any

from backend.app.core.errors import ToolError
from models.adapters import ChatMessage
from tools.base import BaseTool, Permission, ToolContext, ToolResult, ValidationReport
from tools.registry import register

DEFAULT_TARGET_SLIDES = 10


@register
class SlidesTool(BaseTool):
    name = "slides"
    description = "Generate a real PowerPoint (.pptx) deck with titled slides, concise bullets, speaker notes and consistent formatting."
    category = "presentations"
    permissions = [Permission.READ, Permission.WRITE, Permission.FILES, Permission.MODEL]
    cost_estimate = "model-tokens"
    input_schema = {
        "type": "object",
        "properties": {
            "title": {"type": "string"},
            "prompt": {"type": "string", "description": "Topic/brief used to draft the deck"},
            "subtitle": {"type": "string"},
            "audience": {"type": "string"},
            "target_slides": {"type": "integer", "default": 10, "minimum": 8, "maximum": 12},
            "slides": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "title": {"type": "string"},
                        "bullets": {"type": "array", "items": {"type": "string"}},
                        "notes": {"type": "string"},
                    },
                },
            },
            "theme": {"type": "string", "enum": ["default", "dark", "light"], "default": "default"},
            "model": {"type": "string"},
        },
        "required": ["title"],
    }
    output_schema = {
        "type": "object",
        "properties": {"title": {"type": "string"}, "slide_count": {"type": "integer"}, "artifacts": {"type": "array"}},
    }

    def availability(self, ctx: ToolContext | None = None) -> str:
        return "AVAILABLE"

    def execute(self, ctx: ToolContext, payload: dict[str, Any]) -> ToolResult:
        title = str(self.require(payload, "title")).strip()
        if not title:
            raise ToolError("Title must not be empty")

        slides = payload.get("slides") or []
        synthesis = "explicit"
        note = ""
        if not slides:
            if not payload.get("prompt"):
                raise ToolError("Provide either 'slides' or a 'prompt' to draft from")
            slides, synthesis, note = self._draft_slides(payload, ctx, title)

        slides = self._normalise(slides, title)
        deck = {
            "title": title,
            "subtitle": payload.get("subtitle", ""),
            "audience": payload.get("audience", ""),
            "theme": payload.get("theme", "default"),
            "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(),
            "synthesis": synthesis,
            "synthesis_note": note,
            "slides": slides,
        }

        pptx_bytes = build_pptx(deck)
        stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y%m%d-%H%M%S")
        base = re.sub(r"[^A-Za-z0-9._-]+", "-", title).strip("-").lower() or "deck"
        artifact = ctx.save_artifact(
            f"{base}-{stamp}.pptx",
            pptx_bytes,
            "presentation",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            "presentations",
            {"title": title, "slide_count": len(slides)},
        )
        outline = ctx.save_artifact(
            f"{base}-{stamp}-outline.md",
            render_outline(deck),
            "presentation",
            "text/markdown",
            "presentations",
            {"title": title, "role": "outline"},
        )

        return ToolResult(
            status="SUCCESS",
            summary=f"Deck '{title}' generated with {len(slides)} slides.",
            data={"title": title, "slide_count": len(slides), "slides": slides, "deck": deck, "synthesis": synthesis, "synthesis_note": note},
            artifacts=[artifact, outline],
        )

    # -------------------------------------------------------------- helpers
    def _draft_slides(self, payload: dict[str, Any], ctx: ToolContext, title: str) -> tuple[list[dict[str, Any]], str, str]:
        registry = getattr(ctx, "model_registry", None)
        if registry is None:
            from models.registry import registry as global_registry

            registry = global_registry
        target = max(8, min(12, int(payload.get("target_slides", DEFAULT_TARGET_SLIDES))))

        if _chat_usable(registry):
            prompt = (
                f"Create a {target}-slide presentation deck as JSON. Return ONLY a JSON array where each item is "
                '{"title": string, "bullets": [string, ...] (max 5 concise bullets, <=12 words each), "notes": string (2-3 sentence speaker notes)}. '
                "The FIRST slide must be a title slide (its 'title' is the deck title and bullets may hold subtitle info). "
                f"The LAST slide must be a closing/next-steps slide. Deck title: {title}. Audience: {payload.get('audience', 'general')}. "
                f"Brief: {payload.get('prompt')}"
            )
            try:
                completion = registry.complete([ChatMessage("user", prompt)], model=payload.get("model"), max_tokens=3000)
                text = completion.text.strip()
                start, end = text.find("["), text.rfind("]")
                if start != -1 and end != -1:
                    parsed = json.loads(text[start : end + 1])
                    if isinstance(parsed, list) and parsed:
                        return parsed, "llm", ""
            except Exception:  # noqa: BLE001 - deterministic fallback is intended
                pass

        from tools.composition import draft_slides

        return draft_slides(title, str(payload.get("prompt", "")), target), "deterministic", "Model output was unavailable; the deck was composed deterministically from the brief."

    @staticmethod
    def _normalise(slides: list[Any], title: str) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        for item in slides:
            if isinstance(item, str):
                out.append({"title": item, "bullets": [], "notes": ""})
                continue
            if not isinstance(item, dict):
                continue
            slide_title = str(item.get("title") or f"Slide {len(out) + 1}")
            bullets = [str(b) for b in (item.get("bullets") or [])][:6]
            out.append({"title": slide_title, "bullets": bullets, "notes": str(item.get("notes", ""))})
        if not out:
            out = [{"title": title, "bullets": [], "notes": ""}]
        return out

    def validate(self, ctx: ToolContext, result: ToolResult) -> ValidationReport:
        checks = [{"name": "status", "ok": result.status in ("SUCCESS", "UNAVAILABLE"), "detail": result.status}]
        if result.status == "SUCCESS":
            import os
            import zipfile

            checks.append({"name": "slide_count_8_12", "ok": 8 <= result.data.get("slide_count", 0) <= 12, "detail": str(result.data.get("slide_count"))})
            pptx = next((a for a in result.artifacts if a["name"].endswith(".pptx")), None)
            if pptx:
                path = pptx["storage_path"]
                exists = os.path.exists(path) and os.path.getsize(path) > 0
                checks.append({"name": "pptx_exists", "ok": exists, "detail": f"{pptx.get('size', 0)} bytes"})
                if exists:
                    checks.append({"name": "pptx_is_zip", "ok": zipfile.is_zipfile(path), "detail": "OOXML container"})
                    try:
                        from pptx import Presentation

                        prs = Presentation(path)
                        n = len(prs.slides)
                        checks.append({"name": "pptx_readable", "ok": n == result.data.get("slide_count"), "detail": f"{n} slides on disk"})
                    except Exception as exc:  # noqa: BLE001
                        checks.append({"name": "pptx_readable", "ok": False, "detail": str(exc)})
            else:
                checks.append({"name": "pptx_present", "ok": False, "detail": "missing .pptx artifact"})
        return ValidationReport(valid=all(c["ok"] for c in checks), message="slides validation")


# --------------------------------------------------------------------------- #
# Rendering
# --------------------------------------------------------------------------- #
def build_pptx(deck: dict[str, Any]) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    theme = deck.get("theme", "default")
    if theme == "dark":
        title_bg, title_fg, body_fg = RGBColor(0x12, 0x18, 0x2B), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xE5, 0xE7, 0xEB)
    elif theme == "light":
        title_bg, title_fg, body_fg = RGBColor(0xF4, 0xF6, 0xFA), RGBColor(0x11, 0x18, 0x27), RGBColor(0x1F, 0x29, 0x37)
    else:
        title_bg, title_fg, body_fg = RGBColor(0x0F, 0x3D, 0x75), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x1F, 0x29, 0x37)

    slides = deck.get("slides", [])
    total = len(slides)

    for idx, slide_data in enumerate(slides):
        is_first = idx == 0
        layout = prs.slide_layouts[0] if is_first else prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)

        title_placeholder = slide.shapes.title
        title_placeholder.text = slide_data.get("title", "")
        for para in title_placeholder.text_frame.paragraphs:
            for run in para.runs:
                run.font.size = Pt(40 if is_first else 30)
                run.font.bold = True
                run.font.color.rgb = title_fg if is_first else body_fg

        if is_first:
            # Style title slide background.
            from pptx.enum.shapes import MSO_SHAPE

            bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
            bg.fill.solid()
            bg.fill.fore_color.rgb = title_bg
            bg.line.fill.background()
            slide.shapes._spTree.remove(bg._element)
            slide.shapes._spTree.insert(2, bg._element)

        if len(slide.placeholders) > 1:
            body = slide.placeholders[1]
            tf = body.text_frame
            tf.clear()
            bullets = slide_data.get("bullets") or []
            if is_first:
                subtitle = deck.get("subtitle") or ""
                intro = [x for x in [subtitle, deck.get("audience") and f"For: {deck['audience']}"] if x]
                bullets = intro + bullets
            for i, bullet in enumerate(bullets):
                para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                para.text = f"• {bullet}"
                for run in para.runs:
                    run.font.size = Pt(20 if not is_first else 18)
                    run.font.color.rgb = title_fg if is_first else body_fg

        notes = slide_data.get("notes")
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

        # Footer with slide number.
        footer = slide.shapes.add_textbox(Inches(0.4), Inches(7.0), Inches(3), Inches(0.3))
        fpara = footer.text_frame.paragraphs[0]
        fpara.text = f"LocalAI Workspace · {idx + 1}/{total}"
        for run in fpara.runs:
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def render_outline(deck: dict[str, Any]) -> str:
    lines = [f"# {deck['title']}", ""]
    if deck.get("subtitle"):
        lines += [f"_{deck['subtitle']}_", ""]
    if deck.get("synthesis") == "deterministic":
        lines += ["> Composed deterministically from the brief (no language model available).", ""]
    for i, slide in enumerate(deck.get("slides", []), start=1):
        lines.append(f"## Slide {i}: {slide.get('title', '')}")
        for bullet in slide.get("bullets") or []:
            lines.append(f"- {bullet}")
        if slide.get("notes"):
            lines.append(f"  > Notes: {slide['notes']}")
        lines.append("")
    return "\n".join(lines)


def _chat_usable(registry: Any) -> bool:
    if registry is None:
        return False
    checker = getattr(registry, "chat_available", None)
    if callable(checker):
        try:
            return bool(checker())
        except Exception:  # noqa: BLE001
            return False
    try:
        return bool(registry.chat_models())
    except Exception:  # noqa: BLE001
        return False


__all__ = ["SlidesTool", "build_pptx", "render_outline"]
